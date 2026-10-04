"""有界 Excel 读取和可追溯的近一年股票交易行为统计（不调用模型）。"""
from __future__ import annotations

import base64
import binascii
import hashlib
import math
import re
from collections import Counter, defaultdict, deque
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from backend.app.models.trade_history import MAX_FILE_BYTES, TradeHistoryUpload

MAX_ROWS = 20_000
MAX_COLS = 60
FIELD_LABELS = {"date": "成交日期", "code": "证券代码", "side": "买卖方向",
                "quantity": "成交数量", "price": "成交价格", "name": "证券名称",
                "account": "账户标识", "currency": "币种", "asset_type": "证券类型",
                "time": "成交时间", "trade_id": "成交编号"}
REQUIRED = ("date", "code", "side", "quantity", "price")
ALIASES = {
    "date": ("成交日期", "交易日期", "发生日期", "业务日期", "日期", "date", "tradedate"),
    "code": ("证券代码", "股票代码", "代码", "stockcode", "symbol", "code"),
    "side": ("买卖方向", "交易方向", "买卖标志", "操作", "业务名称", "买卖类别", "交易类型", "证券买卖", "side", "direction"),
    "quantity": ("成交数量", "成交股数", "交易数量", "发生数量", "数量", "quantity", "volume", "shares"),
    "price": ("成交价格", "成交均价", "成交价", "交易价格", "价格", "price"),
    "name": ("证券名称", "股票名称", "名称", "name"),
    "account": ("资金账号", "资金账户", "证券账号", "证券账户", "股东账号", "股东账户", "客户号", "客户名称", "用户", "account", "userid"),
    "currency": ("币种", "币别", "currency"),
    "asset_type": ("证券类型", "资产类型", "品种类型", "assettype"),
    "time": ("成交时间", "交易时间", "time"),
    "trade_id": ("成交编号", "成交序号", "tradeid"),
}
BUY = {"买", "买入", "证券买入", "股票买入", "普通买入", "买入成交", "buy", "b"}
SELL = {"卖", "卖出", "证券卖出", "股票卖出", "普通卖出", "卖出成交", "sell", "s"}


def _key(value) -> str:
    return re.sub(r"[\s_（）()\[\]]", "", str(value or "")).casefold()


def _text(value) -> str:
    return str(value if value is not None else "").strip()[:160]


def _decode(request: TradeHistoryUpload) -> bytes:
    if Path(request.file_name).suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError("请选择 .xlsx 或 .xls 格式的 Excel 文件。")
    try:
        raw = base64.b64decode(request.content_base64, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ValueError("文件编码无效，请重新上传 Excel。") from exc
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError("Excel 文件不能为空，且不能超过 5 MB。")
    return raw


@contextmanager
def _workbook(request: TradeHistoryUpload):
    raw = _decode(request)
    book = None
    try:
        if request.file_name.lower().endswith(".xlsx"):
            with ZipFile(BytesIO(raw)) as archive:
                if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                    raise ValueError("Excel 解压后过大，请只导出交易明细工作表。")
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            names = book.sheetnames
            def rows(name):
                sheet = book[name]
                if (sheet.max_column or 0) > MAX_COLS or (sheet.max_row or 0) > MAX_ROWS + 30:
                    raise ValueError("工作表最多支持 20,000 条交易和 60 列，请删除空白格式区域或拆分文件。")
                return sheet.iter_rows(values_only=True)
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=raw, on_demand=True)
            names = book.sheet_names()
            def rows(name):
                sheet = book.sheet_by_name(name)
                if sheet.ncols > MAX_COLS or sheet.nrows > MAX_ROWS + 30:
                    raise ValueError("工作表最多支持 20,000 条交易和 60 列。")
                for n in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(n):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            values.append("#EXCEL_ERROR")
                        else:
                            values.append(cell.value)
                    yield tuple(values)
        if not names or len(names) > 20:
            raise ValueError("Excel 必须包含 1 至 20 个工作表。")
        yield raw, names, rows
    except ImportError as exc:
        raise ValueError("服务缺少 Excel 读取依赖，请按 requirements.txt 安装后重试。") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 Excel，文件可能损坏、加密或格式与扩展名不符。") from exc
    finally:
        if book is not None:
            if hasattr(book, "close"):
                book.close()
            elif hasattr(book, "release_resources"):
                book.release_resources()


def _headers(values):
    headers = []
    for i, v in enumerate(values):
        label = _text(v) or f"未命名列{i + 1}"
        if label in headers:
            label = f"{label}（第{i + 1}列）"
        headers.append(label)
    return headers


def _suggest(headers):
    return {field: next(h for h in headers if _key(h) in aliases)
            for field, aliases in ALIASES.items() if any(_key(h) in aliases for h in headers)}


def _header_row(rows, specified=None, expected=()):
    candidates = []
    for index, values in enumerate(rows, 1):
        headers = _headers(values)
        suggestions = _suggest(headers)
        score = max(sum(k in suggestions for k in REQUIRED), sum(h in headers for h in expected))
        candidates.append((score, index, headers, suggestions))
        if index == specified:
            return index, headers, suggestions
        if index >= 30:
            break
    if specified:
        raise ValueError("指定表头行不存在。")
    if not candidates:
        return 1, [], {}
    _, index, headers, suggestions = max(candidates, key=lambda v: (v[0], -v[1]))
    return index, headers, suggestions


def inspect_trade_workbook(request: TradeHistoryUpload) -> dict:
    with _workbook(request) as (_, names, rows):
        name = request.sheet_name or names[0]
        if name not in names:
            raise ValueError("所选工作表不存在。")
        # 默认优先识别交易工作表，避免把封面或说明页当作数据。
        if request.sheet_name is None:
            for candidate in names:
                _, _, suggested = _header_row(rows(candidate))
                if all(k in suggested for k in REQUIRED):
                    name = candidate
                    break
        header, columns, suggested = _header_row(rows(name), request.header_row)
        preview = []
        # 预览仅输出交易字段；不回传账户姓名等无关原始列。
        for n, values in enumerate(rows(name), 1):
            if n <= header:
                continue
            preview.append({FIELD_LABELS[k]: _text(values[columns.index(v)])
                            for k, v in suggested.items() if k not in {"account", "trade_id"}})
            if len(preview) >= 5:
                break
        return {"sheets": names, "sheet_name": name, "header_row": header,
                "columns": columns, "suggested_columns": suggested, "preview": preview}


def _number(value):
    if isinstance(value, bool):
        raise ValueError("数值无效")
    number = float(re.sub(r"[,，\s]", "", str(value)))
    if not math.isfinite(number) or abs(number) > 1e12:
        raise ValueError("数值无效")
    return number


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if re.fullmatch(r"\d{8}", text):
        return datetime.strptime(text, "%Y%m%d").date()
    text = text.replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
    day = text.split(" ")[0].split("T")[0]
    try:
        return date.fromisoformat(day)
    except ValueError:
        return datetime.strptime(day, "%Y-%m-%d").date()


def _code(value):
    if isinstance(value, bool):
        raise ValueError("证券代码无效")
    if isinstance(value, (float, int)):
        if not math.isfinite(value) or int(value) != value:
            raise ValueError("证券代码无效")
        text = str(int(value)).zfill(6)
    else:
        text = str(value).strip().upper()
        if re.fullmatch(r"\d{1,6}\.0", text):
            text = text[:-2].zfill(6)
        elif re.fullmatch(r"\d{1,4}", text):
            text = text.zfill(6)
    if not re.fullmatch(r"(?:\d{5,6}(?:\.(?:SH|SZ|BJ|HK))?|[A-Z]{1,8}(?:\.[A-Z])?)", text):
        raise ValueError("证券代码无效")
    return re.sub(r"\.(SH|SZ|BJ)$", "", text)


def _time(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.time().isoformat()
    if isinstance(value, time):
        return value.isoformat()
    text = str(value).strip()
    if re.fullmatch(r"\d{6}", text):
        return datetime.strptime(text, "%H%M%S").time().isoformat()
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2})(?:\.\d+)?)?", text)
    if not match:
        raise ValueError("成交时间无效")
    return time(int(match[1]), int(match[2]), int(match[3] or 0)).isoformat()


def analyze_trade_history(request: TradeHistoryUpload, *, today: date | None = None) -> dict:
    today = today or datetime.now(timezone(timedelta(hours=8))).date()
    end = request.as_of or today
    if end > today:
        raise ValueError("分析截止日不能晚于今天。")
    try:
        anniversary = end.replace(year=end.year - 1)
    except ValueError:
        anniversary = end.replace(year=end.year - 1, day=28)
    start = anniversary + timedelta(days=1)
    excluded = Counter()
    issues = []
    trades = []
    accounts = set()
    identities = {}
    with _workbook(request) as (raw, names, rows):
        name = request.sheet_name
        if name is None:
            name = names[0]
            expected = [request.columns[k] for k in REQUIRED if request.columns.get(k)]
            for candidate in names:
                _, candidate_headers, candidate_suggestions = _header_row(rows(candidate), expected=expected)
                if all(request.columns.get(k, candidate_suggestions.get(k)) in candidate_headers for k in REQUIRED):
                    name = candidate
                    break
        if name not in names:
            raise ValueError("所选工作表不存在。")
        header, headers, suggested = _header_row(rows(name), request.header_row,
                                                 [request.columns[k] for k in REQUIRED if request.columns.get(k)])
        columns = {k: v for k, v in {**suggested, **request.columns}.items() if v}
        missing = [FIELD_LABELS[k] for k in REQUIRED if not columns.get(k)]
        if missing:
            raise ValueError("请匹配必填列：" + "、".join(missing))
        if any(v not in headers for v in columns.values()):
            raise ValueError("列名已变化，请重新读取文件并匹配字段。")
        if len(set(columns.values())) != len(columns):
            raise ValueError("每个交易字段需匹配不同的 Excel 列。")
        indices = {k: headers.index(v) for k, v in columns.items()}
        total_rows = 0
        for n, values in enumerate(rows(name), 1):
            if n <= header or not any(v not in (None, "") for v in values):
                continue
            total_rows += 1
            if total_rows > MAX_ROWS:
                raise ValueError("交易记录超过 20,000 条，请拆分文件。")
            data = {k: values[i] if i < len(values) else None for k, i in indices.items()}
            if data.get("account") not in (None, ""):
                accounts.add(_text(data["account"]))
            try:
                d = _date(data["date"])
                if d > end:
                    excluded["截止日之后"] += 1
                    continue
                if d < start:
                    excluded["一年窗口之外"] += 1
                    continue
                side_text = _key(data["side"])
                if side_text not in BUY | SELL:
                    excluded["非普通股票买卖"] += 1
                    continue
                if data.get("asset_type") and _key(data["asset_type"]) not in {"股票", "a股", "b股", "港股", "stock", "equity"}:
                    excluded["非股票品种"] += 1
                    continue
                if data.get("currency") and _key(data["currency"]) not in {"人民币", "cny", "rmb", "元"}:
                    excluded["非人民币交易"] += 1
                    continue
                code = _code(data["code"])
                quantity, price = _number(data["quantity"]), _number(data["price"])
                if side_text in SELL:
                    quantity = abs(quantity)
                if quantity <= 0 or price <= 0 or quantity * price > 1e14:
                    raise ValueError("数量和价格必须为正数且在合理范围内")
                time_value = data.get("time")
                if not time_value and isinstance(data["date"], datetime):
                    time_value = data["date"].time()
                elif not time_value and re.search(r"[T ]\d{1,2}:\d{2}", str(data["date"])):
                    time_value = re.split(r"[T ]", str(data["date"]), maxsplit=1)[1]
                trade = {"date": d.isoformat(), "code": code, "name": _text(data.get("name")) or code,
                         "side": "buy" if side_text in BUY else "sell", "quantity": quantity,
                         "price": price, "amount": quantity * price, "row": n, "time": _time(time_value)}
                identity = _text(data.get("trade_id"))
                if identity:
                    signature = (trade["date"], code, trade["side"], quantity, price, trade["time"])
                    if identity in identities:
                        if identities[identity] != signature:
                            raise ValueError("同一成交编号对应不同记录，请核对原文件")
                        excluded["重复成交编号"] += 1
                        continue
                    identities[identity] = signature
                trades.append(trade)
            except (ValueError, TypeError, OverflowError):
                excluded["字段无效"] += 1
                if len(issues) < 20:
                    issues.append({"row": n, "reason": "日期、代码、数量、价格、币种或成交编号无效，请核对此行。"})
        if len(accounts) > 1:
            raise ValueError("检测到多个账户或用户，请先保留同一用户的交易记录后再上传。")
        if not trades:
            raise ValueError("近一年没有有效的普通股票买卖记录，请检查截止日、列映射和交易类型。")
        source = {"file_name": Path(request.file_name).name, "sha256": hashlib.sha256(raw).hexdigest(),
                  "sheet_name": name, "header_row": header, "columns": {k: v for k, v in columns.items() if k != "account"}}

    trades.sort(key=lambda t: (t["date"], t["time"], t["row"]))
    stock_totals = defaultdict(lambda: {"buy_amount": 0.0, "sell_amount": 0.0, "count": 0})
    monthly = defaultdict(lambda: {"buy_count": 0, "sell_count": 0, "buy_amount": 0.0, "sell_amount": 0.0})
    lots = defaultdict(deque)
    held_days = matched_qty = gross_pnl = 0.0
    matched_sell_count = unmatched_sell_count = 0
    matched_sources = []
    for t in trades:
        s = stock_totals[t["code"]]
        s["name"] = t["name"]
        s[t["side"] + "_amount"] += t["amount"]
        s["count"] += 1
        m = monthly[t["date"][:7]]
        m[t["side"] + "_count"] += 1
        m[t["side"] + "_amount"] += t["amount"]
        d = date.fromisoformat(t["date"])
        if t["side"] == "buy":
            lots[t["code"]].append([t["quantity"], t["price"], d, t["row"]])
            continue
        remaining = t["quantity"]
        trade_pnl = 0.0
        while remaining > 1e-8 and lots[t["code"]]:
            lot = lots[t["code"]][0]
            qty = min(remaining, lot[0])
            days = (d - lot[2]).days
            held_days += days * qty
            matched_qty += qty
            trade_pnl += (t["price"] - lot[1]) * qty
            if len(matched_sources) < 50:
                matched_sources.append({"code": t["code"], "buy_row": lot[3], "sell_row": t["row"], "quantity": qty, "days": days})
            lot[0] -= qty
            remaining -= qty
            if lot[0] < 1e-8:
                lots[t["code"]].popleft()
        if remaining > 1e-8:
            unmatched_sell_count += 1
        else:
            matched_sell_count += 1
            gross_pnl += trade_pnl

    buy_count = sum(t["side"] == "buy" for t in trades)
    buy_amount = sum(t["amount"] for t in trades if t["side"] == "buy")
    sell_amount = sum(t["amount"] for t in trades if t["side"] == "sell")
    turnover = buy_amount + sell_amount
    stocks = [{"code": code, **s, "turnover_share": (s["buy_amount"] + s["sell_amount"]) / turnover}
              for code, s in stock_totals.items()]
    stocks.sort(key=lambda s: s["turnover_share"], reverse=True)
    months = []
    month = start.replace(day=1)
    while month <= end:
        label = month.strftime("%Y-%m")
        months.append({"month": label, **monthly[label]})
        month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
    active_months = sum(m["buy_count"] + m["sell_count"] > 0 for m in months)
    avg_days = held_days / matched_qty if matched_qty else None
    notes = [f"在 {start} 至 {end} 的一年窗口内，有效成交 {len(trades)} 笔，涉及 {len(stocks)} 只股票。",
             f"有交易的月份为 {active_months} 个；买入 {buy_count} 笔，卖出 {len(trades) - buy_count} 笔。",
             f"交易额最高的股票为 {stocks[0]['name']}（{stocks[0]['code']}），占双边成交金额 {stocks[0]['turnover_share']:.1%}；此比例反映交易集中度。"]
    if avg_days is not None:
        notes.append(f"可匹配买卖按先进先出、成交股数加权的平均持有时长为 {avg_days:.1f} 个日历日。")
    limitations = ["历史交易行为不能确定风险承受能力、未来投资期限、流动性需求或收益目标，风险等级仍由确认的问卷决定。",
                   "成交金额按成交数量×成交价格计算，以人民币元计；未计手续费、税费、分红、拆并股及外部转入转出。",
                   "交易额占比不是当前持仓权重；缺少资产余额与估值，不能计算换手率、年化收益或最大回撤。",
                   "持有时长仅计算一年窗口内可按先进先出匹配的股数；期末未卖出部分不计入平均值，盈利仅汇总完全匹配的卖出笔数。",
                   "未提供成交时间时，同日成交按原表行序匹配；没有成交编号时，相同内容仍按独立成交保留。"]
    if unmatched_sell_count:
        limitations.append(f"有 {unmatched_sell_count} 笔卖出无法完整匹配期初成本，未纳入已匹配卖出毛损益。")
    if "currency" not in columns:
        limitations.append("文件未提供币种，按人民币股票记录统计；请确认没有混入外币交易。")
    if "asset_type" not in columns:
        limitations.append("文件未提供证券类型，按用户上传的股票记录统计；请确认没有混入基金、债券或其他品种。")
    if excluded:
        limitations.append(f"共排除 {sum(excluded.values())} 行，具体原因可在数据核验中查看。")
    return {"schema_version": "stock-trades-v1", "source": source, "window_start": start.isoformat(),
            "window_end": end.isoformat(), "first_trade_date": trades[0]["date"], "last_trade_date": trades[-1]["date"],
            "total_rows": total_rows, "trade_count": len(trades), "buy_count": buy_count,
            "sell_count": len(trades) - buy_count, "security_count": len(stocks),
            "active_days": len({t["date"] for t in trades}), "active_months": active_months,
            "buy_amount": round(buy_amount, 2), "sell_amount": round(sell_amount, 2),
            "average_matched_holding_days": round(avg_days, 2) if avg_days is not None else None,
            "matched_sell_count": matched_sell_count, "unmatched_sell_count": unmatched_sell_count,
            "matched_sell_gross_pnl": round(gross_pnl, 2) if matched_sell_count else None,
            "monthly": months, "top_securities": stocks[:10], "behavioral_notes": notes,
            "limitations": limitations, "excluded_rows": dict(excluded), "row_issues": issues,
            "matching_evidence": matched_sources, "preview": [{k: v for k, v in t.items() if k != "time"} for t in trades[:20]],
            "confirmed": False}
