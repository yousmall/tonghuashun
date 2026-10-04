"""交易画像：真实工作簿解析、统计边界、账号保存和页面交互。"""
import base64
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from backend.app import main
from backend.app.database import Database
from backend.app.models.trade_history import TradeHistoryUpload
from backend.app.services.trade_history import analyze_trade_history, inspect_trade_workbook
from backend.app.session_pool import SessionThreadPool

HEADERS = ["成交日期", "证券代码", "买卖方向", "成交数量", "成交价格", "证券名称"]
ROWS = [
    ["2026-01-01", 1, "买入", 100, 10, "示例一"],
    ["2026-01-11", "000001.SZ", "卖出", -40, 12, "示例一"],
    ["2026-01-21", "000001", "卖出", 60, 11, "示例一"],
    ["2026-02-01", "600000", "卖出", 100, 8, "示例二"],
]


def upload(rows=ROWS, headers=HEADERS, **kwargs):
    book = Workbook()
    book.active.title = "说明"
    book.active.append(["券商导出说明"])
    sheet = book.create_sheet("成交明细")
    sheet.append(["股票成交记录"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    return TradeHistoryUpload(file_name="交易.xlsx", content_base64=base64.b64encode(stream.getvalue()).decode(),
                              as_of=date(2026, 10, 4), **kwargs)


def test_inspect_selects_trade_sheet_and_header_and_preserves_identifier():
    result = inspect_trade_workbook(upload())
    assert result["sheet_name"] == "成交明细" and result["header_row"] == 2
    assert result["suggested_columns"]["date"] == "成交日期"
    analysis = analyze_trade_history(upload())
    assert analysis["preview"][0]["code"] == "000001"


def test_real_legacy_xls_and_excel_date_cells():
    raw = (Path(__file__).parent / "fixtures" / "trade_history.xls").read_bytes()
    data = TradeHistoryUpload(file_name="交易.xls", content_base64=base64.b64encode(raw).decode(), as_of=date(2026, 10, 4))
    assert inspect_trade_workbook(data)["header_row"] == 1
    result = analyze_trade_history(data)
    assert result["trade_count"] == 2 and result["average_matched_holding_days"] == 10
    assert result["matched_sell_gross_pnl"] == 200


def test_leap_window_and_unpadded_chinese_date():
    data = upload([["2024年2月1日", 1, "买入", 100, 10, "示例"]]).model_copy(update={"as_of": date(2024, 2, 29)})
    result = analyze_trade_history(data)
    assert result["window_start"] == "2023-03-01"
    assert result["first_trade_date"] == "2024-02-01"


def test_row_column_and_zip_expansion_limits():
    book = Workbook()
    book.active.cell(row=20_031, column=1, value="oversize")
    stream = BytesIO()
    book.save(stream)
    data = upload().model_copy(update={"content_base64": base64.b64encode(stream.getvalue()).decode()})
    with pytest.raises(ValueError, match="20,000"):
        inspect_trade_workbook(data)
    book = Workbook()
    book.active.cell(row=1, column=61, value="oversize")
    stream = BytesIO()
    book.save(stream)
    data.content_base64 = base64.b64encode(stream.getvalue()).decode()
    with pytest.raises(ValueError, match="60"):
        inspect_trade_workbook(data)
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as zipped:
        zipped.writestr("oversize.xml", b"0" * (81 * 1024 * 1024))
    data.content_base64 = base64.b64encode(stream.getvalue()).decode()
    with pytest.raises(ValueError, match="解压后过大"):
        inspect_trade_workbook(data)


def test_fifo_partial_sells_unmatched_opening_lots_and_amount_concentration():
    result = analyze_trade_history(upload())
    assert result["trade_count"] == 4 and result["security_count"] == 2
    assert result["average_matched_holding_days"] == 16
    assert result["matched_sell_count"] == 2 and result["unmatched_sell_count"] == 1
    assert result["matched_sell_gross_pnl"] == 140
    assert result["buy_amount"] == 1000 and result["sell_amount"] == 1940
    assert result["top_securities"][0]["turnover_share"] == pytest.approx(2140 / 2940)
    assert sum(m["buy_count"] + m["sell_count"] for m in result["monthly"]) == 4
    assert result["confirmed"] is False
    assert any("最大回撤" in s for s in result["limitations"])
    assert result["matching_evidence"][0]["buy_row"] == 3


def test_window_excludes_old_future_and_bad_rows_without_inventing_metrics():
    rows = [["2025-10-04", 1, "买入", 100, 10, "旧记录"],
            ["2025-10-05", 1, "买入", 100, 10, "窗口首日"],
            ["2026-10-05", 1, "卖出", 100, 10, "未来记录"],
            ["2026-01-01", 1, "分红", 100, 10, "非成交"],
            ["2026-01-02", 1, "买入", 100, "=10+2", "公式"],
            ["2026-01-03", 1, "买入", -100, 10, "负数量"],
            ["2026-01-04", 1, "卖出", 100, "NaN", "非法数值"]]
    result = analyze_trade_history(upload(rows))
    assert result["trade_count"] == 1
    assert result["window_start"] == "2025-10-05"
    assert result["excluded_rows"] == {"一年窗口之外": 1, "截止日之后": 1, "非普通股票买卖": 1, "字段无效": 3}
    assert result["average_matched_holding_days"] is None
    assert result["matched_sell_gross_pnl"] is None


def test_time_order_handles_unpadded_hours_and_date_time_cells():
    rows = [[datetime(2026, 1, 1, 10, 30), 1, "卖出", 100, 12, "示例"],
            [datetime(2026, 1, 1, 9, 30), 1, "买入", 100, 10, "示例"]]
    result = analyze_trade_history(upload(rows))
    assert result["matched_sell_gross_pnl"] == 200 and result["average_matched_holding_days"] == 0
    rows = [["2026-01-01", 1, "卖出", 100, 12, "示例", "10:30"],
            ["2026-01-01", 1, "买入", 100, 10, "示例", "9:30"]]
    assert analyze_trade_history(upload(rows, HEADERS + ["成交时间"]))["matched_sell_gross_pnl"] == 200


def test_multi_account_rejected_and_preview_does_not_include_account_data():
    rows = [ROW + ["account-a"] for ROW in ROWS]
    data = upload(rows, HEADERS + ["资金账号"])
    assert "account-a" not in str(inspect_trade_workbook(data)["preview"])
    analyze_trade_history(data)
    rows[1][-1] = "account-b"
    with pytest.raises(ValueError, match="多个账户"):
        analyze_trade_history(upload(rows, HEADERS + ["资金账号"]))


def test_custom_columns_missing_columns_and_optional_ignore():
    headers = ["dt", "ticker", "action", "qty", "unit_price", "label", "币种"]
    data = upload([r + ["USD"] for r in ROWS], headers,
                  columns=dict(zip(["date", "code", "side", "quantity", "price", "name", "currency"], headers)))
    with pytest.raises(ValueError, match="没有有效"):
        analyze_trade_history(data)
    data.columns["currency"] = ""
    assert analyze_trade_history(data)["trade_count"] == 4
    with pytest.raises(ValueError, match="必填列"):
        analyze_trade_history(upload(ROWS, headers))
    data.columns["code"] = "dt"
    with pytest.raises(ValueError, match="不同"):
        analyze_trade_history(data)


def test_duplicate_trade_id_deduplicates_but_repeated_fills_without_id_remain():
    rows = [ROWS[0] + ["fill-1"], ROWS[0] + ["fill-1"]]
    result = analyze_trade_history(upload(rows, HEADERS + ["成交编号"]))
    assert result["trade_count"] == 1 and result["excluded_rows"]["重复成交编号"] == 1
    assert analyze_trade_history(upload([ROWS[0], ROWS[0]]))["trade_count"] == 2


@pytest.mark.parametrize("patch, message", [
    ({"content_base64": "%%%"}, "编码"),
    ({"content_base64": base64.b64encode(b"not excel").decode()}, "无法读取"),
    ({"file_name": "trade.csv"}, "格式"),
    ({"sheet_name": "不存在"}, "工作表"),
    ({"as_of": date(2099, 1, 1)}, "晚于今天"),
])
def test_bad_inputs(patch, message):
    with pytest.raises(ValueError, match=message):
        analyze_trade_history(upload().model_copy(update=patch))


def test_api_auth_preview_confirmation_version_isolation_and_questionnaire_preservation(monkeypatch, request, risk_questionnaire_payload):
    database = Database("sqlite+pysqlite:///:memory:", "test-secret-that-is-longer-than-thirty-two-characters")
    database.initialize()
    monkeypatch.setattr(main, "database", database)
    pool = SessionThreadPool()
    request.addfinalizer(pool.shutdown)
    monkeypatch.setattr(main, "session_thread_pool", pool)
    payload = upload().model_dump(mode="json")
    with TestClient(main.app) as client:
        assert client.post("/api/v1/profile/trades/analyze", json=payload).status_code == 401
        login = client.post("/api/v1/auth/register", json={"username": "trader-one", "password": "Strong-trade-123"}).json()
        headers = {"Authorization": f"Bearer {login['access_token']}"}
        assert client.post("/api/v1/profile/trades/inspect", headers=headers, json=payload).status_code == 200
        assert client.post("/api/v1/profile/trades/analyze", headers=headers, json=payload).status_code == 200
        assert database.get_profile(login["user"]["id"]) is None
        saved = client.post("/api/v1/profile/trades/confirm", headers=headers, json=payload)
        assert saved.status_code == 200, saved.text
        profile = saved.json()
        assert profile["version"] == 2 and not profile["confirmed"] and profile["risk_level"] is None
        assert profile["trading_analysis"]["confirmed"] and "preview" not in profile["trading_analysis"]
        assert client.post("/api/v1/profile/trades/confirm", headers=headers, json=payload).status_code == 409
        draft = client.post("/api/v1/profile/assess", headers=headers, json=risk_questionnaire_payload).json()["profile"]
        draft.update(version=2, trading_analysis={"risk_level": "R5", "trade_count": 99999})
        confirmed = client.post("/api/v1/profile/confirm", headers=headers, json={"profile": draft})
        assert confirmed.status_code == 200, confirmed.text
        confirmed = confirmed.json()
        assert confirmed["trading_analysis"] == profile["trading_analysis"] and confirmed["confirmed"]
        payload["expected_version"] = confirmed["version"]
        updated = client.post("/api/v1/profile/trades/confirm", headers=headers, json=payload).json()
        for key in ("risk_answers", "risk_level", "risk_score", "valid_until", "confirmed"):
            assert updated[key] == confirmed[key]
        assert client.get("/api/v1/profile", headers=headers).json()["profile"] == updated
        other = client.post("/api/v1/auth/register", json={"username": "trader-two", "password": "Strong-trade-123"}).json()
        other_headers = {"Authorization": f"Bearer {other['access_token']}"}
        assert not client.get("/api/v1/profile", headers=other_headers).json()["profile"]["trading_analysis"]
        invalid = {**payload, "content_base64": "%%%"}
        assert client.post("/api/v1/profile/trades/confirm", headers=headers, json=invalid).status_code == 422
        assert client.get("/api/v1/profile", headers=headers).json()["profile"] == updated


def test_page_excel_flow_displays_summary_and_saves_without_risk_confirmation():
    # AppTest 尚不支持设置上传控件，用真实 xlsx 字节替换上传返回值，其余控件真实执行。
    content = upload().content_base64
    script = f'''
import base64
from io import BytesIO
from unittest.mock import patch
import streamlit as st
from frontend import streamlit_app as ui
from backend.app.models.trade_history import TradeHistoryUpload
from backend.app.services.trade_history import inspect_trade_workbook, analyze_trade_history
ui.init_session()
file = BytesIO(base64.b64decode({content!r}))
file.name = "交易.xlsx"
def fake_api(base, method, path, payload=None, **kwargs):
    data = TradeHistoryUpload(**payload)
    if path.endswith('/inspect'):
        return inspect_trade_workbook(data)
    analysis = analyze_trade_history(data)
    if path.endswith('/analyze'):
        return analysis
    analysis['confirmed'] = True
    st.session_state.did_save = True
    return dict(st.session_state.profile, trading_analysis=analysis, version=2)
with patch('frontend.trade_profile.st.file_uploader', return_value=file):
    from frontend.trade_profile import render_trade_profile
    render_trade_profile('http://offline', api_request=fake_api)
'''
    app = AppTest.from_string(script).run(timeout=15)
    assert not app.exception
    next(b for b in app.button if b.label == "读取 Excel").click().run(timeout=15)
    assert not app.exception
    assert app.selectbox[0].value == "成交明细"
    next(b for b in app.button if b.label == "分析交易画像").click().run(timeout=15)
    assert not app.exception and app.metric[0].value == "4 笔"
    next(b for b in app.button if b.label == "确认并保存交易画像").click().run(timeout=15)
    assert not app.exception and app.session_state["did_save"]
    assert not app.session_state["profile"]["confirmed"]
    assert app.session_state["profile"]["trading_analysis"]["confirmed"]


def test_logout_clears_upload_and_analysis_state():
    script = '''
import streamlit as st
from frontend import streamlit_app as ui
ui.init_session()
st.session_state.trade_profile_draft = {'private': 'data'}
st.session_state.trade_profile_upload = 'file'
ui.reset_user_session()
'''
    app = AppTest.from_string(script).run()
    assert not app.exception
    assert "trade_profile_draft" not in app.session_state and "trade_profile_upload" not in app.session_state
